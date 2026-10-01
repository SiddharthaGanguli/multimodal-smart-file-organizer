"""Real document fixtures exercise extraction, OCR hand-off, and saved results."""

import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
from reportlab.pdfgen import canvas

from app.extractors import extract_and_store, extract_document
from app.extractors.models import ExtractionResult


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16"])
def test_txt_unicode_and_original_preserved(tmp_path, encoding):
    source = tmp_path / "notes.TXT"
    source.write_bytes("Invoice ₹250\r\nनमस्ते".encode(encoding))
    original = source.read_bytes()
    result = extract_document(source)
    assert result.status == "extracted"
    assert result.text == "Invoice ₹250\nनमस्ते"
    assert result.encoding
    assert source.read_bytes() == original


@pytest.mark.parametrize("content", [b"", b" \n\t"])
def test_empty_txt(tmp_path, content):
    source = tmp_path / "empty.txt"
    source.write_bytes(content)
    assert extract_document(source).status == "empty"


@pytest.mark.parametrize("content", [b"\x80\xff", b"a\x00b"])
def test_invalid_txt_does_not_silently_replace_characters(tmp_path, content):
    source = tmp_path / "bad.txt"
    source.write_bytes(content)
    assert extract_document(source).status == "failed"


def test_docx_paragraph_table_order_and_nested_table(tmp_path):
    source = tmp_path / "invoice.docx"
    document = Document()
    document.add_paragraph("Invoice heading")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Coffee"
    table.cell(0, 1).text = "250"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested detail"
    document.add_paragraph("Thank you")
    document.save(source)
    original = source.read_bytes()
    result = extract_document(source)
    assert result.status == "extracted"
    assert result.text.index("Invoice heading") < result.text.index("Coffee")
    assert result.text.index("250") < result.text.index("Thank you")
    assert "Nested detail" in result.text
    assert [part.location for part in result.parts] == [
        "body/paragraph/1",
        "body/table/2",
        "body/paragraph/3",
    ]
    assert all(part.page_number is None for part in result.parts)
    assert source.read_bytes() == original


def test_empty_docx(tmp_path):
    source = tmp_path / "empty.docx"
    Document().save(source)
    assert extract_document(source).status == "empty"


@pytest.mark.parametrize("location", ["block", "inline", "cell", "nested"])
def test_docx_content_controls_preserve_text(tmp_path, location):
    source = tmp_path / "controlled.docx"
    document = Document()
    document.add_paragraph("Start")
    if location == "cell":
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Revenue: INR 2900"
        element = cell._tc
    else:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("Revenue: INR 2900")
        element = run._r if location == "inline" else paragraph._p
    for _ in range(2 if location == "nested" else 1):
        parent = element.getparent()
        control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        parent.replace(element, control)
        content.append(element)
        control.append(content)
        element = control
    document.add_paragraph("End")
    document.save(source)
    original = source.read_bytes()
    result = extract_document(source)
    assert result.status == "extracted"
    assert result.text == "Start\n\nRevenue: INR 2900\n\nEnd"
    assert source.read_bytes() == original


@pytest.mark.parametrize("vertical", [False, True])
def test_docx_merged_cell_is_extracted_once(tmp_path, vertical):
    source = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    endpoint = table.cell(1, 0) if vertical else table.cell(0, 1)
    table.cell(0, 0).merge(endpoint).text = "Total: INR 900"
    table.cell(1, 1).text = "Approved"
    document.save(source)
    result = extract_document(source)
    assert result.status == "extracted"
    assert result.text.count("Total: INR 900") == 1
    assert result.text.count("Approved") == 1


def test_docx_distinct_cells_with_identical_text_are_preserved(tmp_path):
    source = tmp_path / "repeated.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    for cell in table.rows[0].cells:
        cell.text = "INR 900"
    document.save(source)
    assert extract_document(source).text == "INR 900\tINR 900"


def make_pdf(path: Path, pages: list[str | None], encrypt=None):
    document = canvas.Canvas(str(path), encrypt=encrypt)
    for text in pages:
        if text is None:
            # An image-only page, like a scanned document.
            document.drawInlineImage(Image.new("RGB", (20, 20), "black"), 40, 40)
        elif text:
            document.drawString(72, 720, text)
        document.showPage()
    document.save()


def test_pdf_text_and_page_provenance(tmp_path):
    source = tmp_path / "invoice.pdf"
    make_pdf(source, ["Invoice 123", "Total 250"])
    original = source.read_bytes()
    result = extract_document(source)
    assert result.status == "extracted"
    assert "Invoice 123" in result.text and "Total 250" in result.text
    assert [part.page_number for part in result.parts] == [1, 2]
    assert result.ocr_pages == []
    assert source.read_bytes() == original


@pytest.mark.parametrize(
    "pages,ocr_pages",
    [
        ([None], [1]),
        (["Invoice", None, "Total"], [2]),
        ([""], [1]),
    ],
)
def test_pdf_ocr_detection(tmp_path, pages, ocr_pages):
    source = tmp_path / "scan.pdf"
    make_pdf(source, pages)
    result = extract_document(source)
    assert result.status == "needs_ocr"
    assert result.ocr_pages == ocr_pages
    assert len(result.parts) == len(pages)
    for text in pages:
        if text:
            assert text in result.text


def test_password_protected_pdf(tmp_path):
    source = tmp_path / "locked.pdf"
    make_pdf(source, ["Private invoice"], encrypt="secret")
    result = extract_document(source)
    assert result.status == "failed"
    assert result.text == ""


@pytest.mark.parametrize(
    "large_image,caption,needs_ocr",
    [
        (True, "1", True),
        (False, "Invoice 123", False),
        (
            True,
            (
                "This is a searchable report containing enough actual text to describe the invoice "
                "and its payment details including the total amount of INR 900."
            ),
            False,
        ),
    ],
)
def test_pdf_scan_with_page_number_and_searchable_image_pages(
    tmp_path,
    large_image,
    caption,
    needs_ocr,
):
    source = tmp_path / "image-and-text.pdf"
    document = canvas.Canvas(str(source))
    width, height = (550, 670) if large_image else (30, 30)
    document.drawInlineImage(
        Image.new("RGB", (100, 100), "white"), 30, 80, width=width, height=height
    )
    document.drawString(40, 30, caption)
    document.save()
    result = extract_document(source)
    assert result.status == ("needs_ocr" if needs_ocr else "extracted")
    assert result.ocr_pages == ([1] if needs_ocr else [])
    assert caption in result.text


@pytest.mark.parametrize("suffix", ["pdf", "docx"])
def test_corrupt_documents(tmp_path, suffix):
    source = tmp_path / f"broken.{suffix}"
    source.write_bytes(b"This is not a valid document container.")
    result = extract_document(source)
    assert result.status == "failed"
    assert result.error


def test_missing_and_unsupported_files(tmp_path):
    assert extract_document(tmp_path / "missing.txt").status == "failed"
    result = extract_document(tmp_path / "photo.jpg")
    assert result.status == "failed"
    assert result.error == "Unsupported file type."


@pytest.mark.parametrize("exists", [True, False])
def test_persist_success_and_failure_with_provenance(tmp_path, exists):
    source = tmp_path / "notes.txt"
    if exists:
        source.write_text("Invoice", encoding="utf-8")
    result, destination = extract_and_store(source, tmp_path / "results")
    saved = ExtractionResult.model_validate_json(destination.read_text(encoding="utf-8"))
    assert saved == result
    assert saved.status == ("extracted" if exists else "failed")
    assert saved.source_name == "notes.txt"
    assert saved.extractor == "python"
    assert saved.extractor_version
    assert saved.extracted_at.tzinfo is not None
    _, second = extract_and_store(source, destination.parent)
    assert second != destination
    assert destination.exists()


def test_cli_saves_result_and_reports_failure(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("CLI invoice", encoding="utf-8")
    command = [
        sys.executable,
        "-m",
        "app.extractors",
        str(source),
        "--output-dir",
        str(tmp_path / "results"),
    ]
    process = subprocess.run(command, capture_output=True, text=True, check=False)
    assert process.returncode == 0, process.stderr
    assert "Status: extracted" in process.stdout
    source.unlink()
    process = subprocess.run(command, capture_output=True, text=True, check=False)
    assert process.returncode == 1
    assert "Status: failed" in process.stdout
    assert len(list((tmp_path / "results").glob("*.json"))) == 2


def test_default_storage_and_saved_ocr_status(tmp_path, monkeypatch):
    from app.core.config import get_settings

    source = tmp_path / "mixed.pdf"
    make_pdf(source, ["Invoice", None])
    monkeypatch.setattr(get_settings(), "storage_dir", str(tmp_path / "private"))
    result, destination = extract_and_store(source)
    assert destination.parent == tmp_path / "private" / "extractions"
    saved = ExtractionResult.model_validate_json(destination.read_text(encoding="utf-8"))
    assert saved == result
    assert saved.status == "needs_ocr"
    assert saved.ocr_pages == [2]
    assert saved.extractor == "pdfminer.six"


def test_storage_failure_is_not_reported_as_success(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("Original", encoding="utf-8")
    with pytest.raises(OSError):
        extract_and_store(source, output_dir=source)
    assert source.read_text(encoding="utf-8") == "Original"
