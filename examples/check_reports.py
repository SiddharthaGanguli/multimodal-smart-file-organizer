"""Generate fictional reports and compare extraction against known expected text.

Run from the repository root: python examples/check_reports.py
Requires the project's development dependencies.
"""

import hashlib
import json
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas

from app.extractors import extract_and_store


def check_reports(directory: Path) -> list[dict]:
    directory.mkdir(parents=True, exist_ok=True)
    cases = []

    text = (
        "Monthly expense report\nTravel: INR 1200\nMeals: INR 450\nTotal: INR 1650\nनोट: भुगतान हुआ"
    )
    source = directory / "expenses.txt"
    source.write_text(text, encoding="utf-8")
    cases.append((source, text, "extracted", []))

    source = directory / "sales.docx"
    document = Document()
    document.add_heading("Quarterly sales report", 0)
    document.add_paragraph("Period: January to March 2026")
    table = document.add_table(rows=3, cols=3)
    rows = [
        ("Product", "Units", "Revenue"),
        ("Notebook", "10", "INR 2500"),
        ("Pen", "20", "INR 400"),
    ]
    for row_index, values in enumerate(rows):
        for column, value in enumerate(values):
            table.cell(row_index, column).text = value
    document.add_paragraph("Total revenue: INR 2900")
    document.add_paragraph("Reviewed by: Demo Reviewer")
    document.save(source)
    expected = "\n\n".join(
        [
            "Quarterly sales report",
            "Period: January to March 2026",
            "\n".join("\t".join(row) for row in rows),
            "Total revenue: INR 2900",
            "Reviewed by: Demo Reviewer",
        ]
    )
    cases.append((source, expected, "extracted", []))

    source = directory / "project.pdf"
    pages = [
        ["Project progress report", "Completed tasks: 12", "Pending tasks: 3"],
        ["Budget summary", "Allocated: INR 10000", "Spent: INR 7500", "Balance: INR 2500"],
    ]
    pdf = canvas.Canvas(str(source))
    for lines in pages:
        for index, line in enumerate(lines):
            pdf.drawString(72, 720 - index * 30, line)
        pdf.showPage()
    pdf.save()
    cases.append((source, "\n".join(line for page in pages for line in page), "extracted", []))

    source = directory / "mixed-scan.pdf"
    scan = Image.new("RGB", (900, 500), "white")
    ImageDraw.Draw(scan).text((40, 40), "Scanned invoice: INR 300", fill="black", font_size=32)
    pdf = canvas.Canvas(str(source))
    pdf.drawString(72, 720, "Cover: supporting invoice follows")
    pdf.showPage()
    pdf.drawInlineImage(scan, 50, 200, width=500, height=280)
    pdf.showPage()
    pdf.save()
    cases.append((source, "Cover: supporting invoice follows", "needs_ocr", [2]))

    checks = []
    for source, expected, status, ocr_pages in cases:
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        result, destination = extract_and_store(source, directory / "results")
        # PDF layout introduces whitespace; compare every word and its order.
        matches = (
            result.text.split() == expected.split()
            if source.suffix == ".pdf"
            else result.text == expected
        )
        unchanged = before == hashlib.sha256(source.read_bytes()).hexdigest()
        saved = json.loads(destination.read_text(encoding="utf-8"))
        passed = (
            matches
            and unchanged
            and result.status == status
            and result.ocr_pages == ocr_pages
            and saved["text"] == result.text
        )
        checks.append(
            {
                "source": str(source),
                "passed": passed,
                "text_matches": matches,
                "original_unchanged": unchanged,
                "expected_text": expected,
                "actual_text": result.text,
                "expected_status": status,
                "actual_status": result.status,
                "expected_ocr_pages": ocr_pages,
                "actual_ocr_pages": result.ocr_pages,
                "result_file": str(destination),
            }
        )
    return checks


def main() -> int:
    directory = Path("storage/report-checks")
    checks = check_reports(directory)
    report = directory / "verification.json"
    report.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    for check in checks:
        outcome = "PASS" if check["passed"] else "FAIL"
        print(f"{outcome}: {check['source']} ({check['actual_status']})")
    print(f"Comparison report: {report}")
    return 0 if all(check["passed"] for check in checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
